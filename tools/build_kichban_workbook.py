# -*- coding: utf-8 -*-
"""
build_kichban_workbook.py — từ JSON hội thoại đã phân loại → file .xlsx 5 sheet,
khớp khuôn template/Kangnam_KichBan_CatMi.xlsx.

    pip install openpyxl
    python tools/build_kichban_workbook.py chats/cat-mi.json
    python tools/build_kichban_workbook.py "chats/*.json" -o out/

Generator TỪ CHỐI BUILD nếu dữ liệu sai — xem mục "CỔNG KIỂM TRA" ở cuối file.
Không sửa script để lách lỗi: sửa JSON rồi chạy lại.
"""
import argparse
import glob
import json
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Thiếu openpyxl. Chạy: pip install openpyxl")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# ============================================================ BẢNG MÀU (lấy từ template)
NAVY = "FF1F3864"      # tiêu đề lớn
BLUE = "FF2E5496"      # header bảng
L_BLUE = "FFD9E1F2"    # nền phụ đề
CREAM = "FFFFF2CC"     # nhóm A / cảnh báo
GOLD = "FFBF8F00"      # chữ nhóm A
GREEN_BG = "FFE2EFDA"  # nhóm B / đạt
GREEN = "FF548235"
SKY = "FFDDEBF7"       # nhóm C / WARM
RED_BG = "FFFCE4E4"    # HOT / lỗi
RED = "FFC00000"
GREY_BG = "FFF2F2F2"
GREY = "FF808080"
WHITE = "FFFFFFFF"
INK = "FF000000"
MUTE = "FF595959"

F = lambda c: PatternFill("solid", start_color=c, end_color=c)
WRAP_T = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center")

# ============================================================ CỔNG KIỂM TRA — dữ liệu
TU_CAM = [
    "không đau", "không sưng", "cam kết", "trọn đời", "tốt nhất", "duy nhất",
    "số 1", "100%", "an toàn tuyệt đối", "tuyệt đối an toàn", "không biến chứng",
    "không để lại sẹo", "không để lại dấu vết", "khỏi hẳn", "vĩnh viễn",
    "ưu đãi bí mật", "chắc chắn khỏi",
]
# giá cứng trong tin nhắn: "8 triệu", "8.000.000", "8tr", "10tr → 8tr"…
RE_GIA = re.compile(
    r"\d[\d.,]*\s*(?:tr(?:iệu)?\b|đ\b|vnđ\b|vnd\b|k\b)|\d{1,3}(?:\.\d{3}){2,}",
    re.I)
DO_NONG_HOP_LE = ["HOT", "WARM", "COLD", "Đã booking", "Hậu phẫu", "Loại"]
SDT_HOP_LE = ["Có", "Không"]


def diem_tin_hieu(ma):
    """[10110] → 3. Trả về None nếu mã sai khuôn."""
    m = re.fullmatch(r"\[([01]{5})\]", (ma or "").strip())
    return sum(int(x) for x in m.group(1)) if m else None


def do_nong_tu_diem(d):
    return "HOT" if d >= 4 else ("WARM" if d >= 2 else "COLD")


def validate(cfg):
    e = []
    need = ["zone", "nguon", "nhom_dt", "hoi_thoai", "kich_ban", "do_nong", "l_code",
            "van_de", "tin_hieu", "loi_quan_sat", "bien_so"]
    for k in need:
        if k not in cfg:
            e.append("Thiếu khối bắt buộc: '%s'" % k)
    if e:
        return e

    z = cfg["zone"]
    for k in ["ten", "dich_vu", "ten_file"]:
        if not z.get(k):
            e.append("zone.%s trống" % k)

    ma_nhom = [n["ma"] for n in cfg["nhom_dt"]]
    ma_lcode = [l["ma"] for l in cfg["l_code"]]
    ma_kb = [k["ma"] for k in cfg["kich_ban"]]

    # ---- hội thoại
    seen = set()
    for i, h in enumerate(cfg["hoi_thoai"], 1):
        tag = "hoi_thoai[%d]" % i
        for k in ["ten_kh", "ngay", "phu_trach", "so_tin", "sdt", "van_de",
                  "nhom_dt", "l_code", "do_nong", "ghi_chu"]:
            if k not in h:
                e.append("%s: thiếu '%s'" % (tag, k))
        if e and len(e) > 60:
            return e
        if h.get("sdt") not in SDT_HOP_LE:
            e.append("%s: sdt='%s' — chỉ nhận %s" % (tag, h.get("sdt"), SDT_HOP_LE))
        if h.get("do_nong") not in DO_NONG_HOP_LE:
            e.append("%s: do_nong='%s' — chỉ nhận %s" % (tag, h.get("do_nong"), DO_NONG_HOP_LE))
        if h.get("nhom_dt") not in ma_nhom:
            e.append("%s: nhom_dt='%s' chưa khai trong nhom_dt[]" % (tag, h.get("nhom_dt")))
        if h.get("l_code") not in ma_lcode:
            e.append("%s: l_code='%s' chưa khai trong l_code[]" % (tag, h.get("l_code")))
        if h.get("van_de") not in cfg["van_de"]:
            e.append("%s: van_de='%s' chưa khai trong van_de[]" % (tag, h.get("van_de")))

        # CỔNG TÍN HIỆU: mã 5 chữ số và độ nóng phải khớp điểm
        ma = re.match(r"(\[[^\]]*\])", (h.get("ghi_chu") or "").strip())
        if not ma:
            e.append("%s: ghi_chu phải bắt đầu bằng mã 5 tín hiệu dạng [01011]" % tag)
        else:
            d = diem_tin_hieu(ma.group(1))
            if d is None:
                e.append("%s: mã tín hiệu '%s' sai khuôn (cần 5 chữ số 0/1)" % (tag, ma.group(1)))
            elif h.get("do_nong") in ("HOT", "WARM", "COLD"):
                dung = do_nong_tu_diem(d)
                if h["do_nong"] != dung:
                    e.append("%s: %s điểm → phải là %s, đang ghi %s"
                             % (tag, d, dung, h["do_nong"]))

        # CỔNG SĐT ↔ NHÓM: nhóm A = chưa số, nhóm B/C = đã có số
        n = h.get("nhom_dt") or ""
        if n.startswith("A") and h.get("sdt") == "Có":
            e.append("%s: nhóm %s là nhóm CHƯA có số nhưng sdt='Có'" % (tag, n))
        if n[:1] in ("B", "C") and h.get("sdt") == "Không":
            e.append("%s: nhóm %s là nhóm ĐÃ có số nhưng sdt='Không'" % (tag, n))

        key = (h.get("ten_kh"), h.get("ngay"))
        if key in seen:
            e.append("%s: trùng (tên + ngày) với dòng trước — %s" % (tag, key))
        seen.add(key)

    # ---- kịch bản
    for i, k in enumerate(cfg["kich_ban"], 1):
        tag = "kich_ban[%d] %s" % (i, k.get("ma", "?"))
        for f in ["nhom", "ma", "giai_doan", "insight", "l_code", "khi_dung",
                  "noi_dung", "luu_y"]:
            if f not in k:
                e.append("%s: thiếu '%s'" % (tag, f))
        nd = k.get("noi_dung") or []
        if not isinstance(nd, list) or not (1 <= len(nd) <= 3):
            e.append("%s: noi_dung phải là list 1–3 tin" % tag)
            continue
        for j, tin in enumerate(nd, 1):
            low = (tin or "").lower()
            hit = [w for w in TU_CAM if w in low]
            if hit:
                e.append("%s nội dung %d: chứa từ cấm %s" % (tag, j, hit))
            g = RE_GIA.search(tin or "")
            if g:
                e.append("%s nội dung %d: có GIÁ CỨNG '%s' — tin nhắn chỉ nêu khung "
                         "và mời khám, số cụ thể lấy từ bien_so" % (tag, j, g.group(0).strip()))
        low_ly = (k.get("luu_y") or "").lower()
        if not low_ly:
            e.append("%s: luu_y trống — mỗi kịch bản phải khai rào an toàn" % tag)

    # mỗi nhóm đối tượng phải có ít nhất 1 kịch bản
    for n in ma_nhom:
        if n not in ma_kb:
            e.append("nhom_dt '%s' không có kịch bản nào trong kich_ban[]" % n)

    # ---- độ nóng
    for i, d in enumerate(cfg["do_nong"], 1):
        for f in ["muc", "nghia", "tin_hieu", "sla", "nhip", "kb", "kpi"]:
            if f not in d:
                e.append("do_nong[%d]: thiếu '%s'" % (i, f))
        if d.get("muc") not in DO_NONG_HOP_LE:
            e.append("do_nong[%d]: muc='%s' không hợp lệ" % (i, d.get("muc")))

    if len(cfg["tin_hieu"]) != 5:
        e.append("tin_hieu[] phải đúng 5 phần tử (5 tín hiệu chấm độ nóng)")
    return e


# ============================================================ DỰNG SHEET
def _title(ws, text, ncol, size=13):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, text)
    c.font = Font(bold=True, size=size, color=WHITE)
    c.fill = F(NAVY)
    c.alignment = WRAP_C
    ws.row_dimensions[1].height = 30


def _note(ws, row, text, ncol):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row, 1, text)
    c.font = Font(size=9, color=RED)
    c.fill = F(CREAM)
    c.alignment = WRAP_T
    ws.row_dimensions[row].height = 28


def _head(ws, row, cols):
    for i, (name, w) in enumerate(cols, 1):
        c = ws.cell(row, i, name)
        c.font = Font(bold=True, size=10, color=WHITE)
        c.fill = F(BLUE)
        c.alignment = WRAP_C
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 1)


MAU_DO_NONG = {
    "HOT": (RED, RED_BG), "WARM": (GOLD, CREAM), "COLD": (BLUE, SKY),
    "Đã booking": (GREEN, GREEN_BG), "Hậu phẫu": (GREEN, GREEN_BG),
    "Loại": (GREY, GREY_BG),
}
MAU_NHOM = {"A": (GOLD, CREAM), "B": (GREEN, GREEN_BG), "C": (BLUE, SKY)}


def sheet_tong_quan(wb, cfg, agg):
    ws = wb.create_sheet("Tổng quan")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", (3, 34, 60, 30)):
        ws.column_dimensions[col].width = w
    z = cfg["zone"]

    ws.merge_cells("B2:D2")
    c = ws.cell(2, 2, "KANGNAM · %s — PHÂN LOẠI HỘI THOẠI & BỘ KỊCH BẢN TRỰC PAGE"
                % z["dich_vu"].upper())
    c.font, c.fill, c.alignment = Font(bold=True, size=15, color=WHITE), F(NAVY), WRAP_C
    ws.row_dimensions[2].height = 32

    ws.merge_cells("B3:D3")
    c = ws.cell(3, 2, cfg["nguon"])
    c.font, c.fill, c.alignment = Font(size=9, color=MUTE), F(L_BLUE), WRAP_C
    ws.row_dimensions[3].height = 26

    r = 5

    def section(txt):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(r, 2, txt)
        c.font, c.fill, c.alignment = Font(bold=True, size=11, color=WHITE), F(BLUE), WRAP_T
        ws.row_dimensions[r].height = 22
        r += 1

    def kv(k, v, note="", kfill=GREY_BG, kcolor=INK):
        nonlocal r
        a = ws.cell(r, 2, k)
        a.font, a.fill, a.alignment = Font(bold=True, size=10, color=kcolor), F(kfill), WRAP_T
        b = ws.cell(r, 3, v)
        b.font, b.alignment = Font(bold=True, size=10, color=NAVY), WRAP_T
        d = ws.cell(r, 4, note)
        d.font, d.alignment = Font(size=9), WRAP_T
        r += 1

    def kv2(k, v, kfill=WHITE, kcolor=INK, kbold=True):
        """2 cột: nhãn + mô tả dài (gộp C:D)."""
        nonlocal r
        a = ws.cell(r, 2, k)
        a.font, a.fill, a.alignment = Font(bold=kbold, size=10, color=kcolor), F(kfill), WRAP_T
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        b = ws.cell(r, 3, v)
        b.font, b.fill, b.alignment = Font(size=9), F(kfill), WRAP_T
        r += 1

    # ① kết quả phân loại — TÍNH TỪ DỮ LIỆU, không nhập tay
    section("1 · KẾT QUẢ PHÂN LOẠI %d HỘI THOẠI" % agg["tong"])
    kv("Tổng hội thoại phân tích", agg["tong"], cfg.get("ghi_chu_tong", ""))
    kv("CÓ để lại SĐT", "%d (%d%%)" % (agg["co_sdt"], agg["pc_co"]),
       "Đưa vào nhóm B — mục tiêu: kết nối nhanh, đẩy lên Booking")
    kv("KHÔNG để lại SĐT", "%d (%d%%)" % (agg["khong_sdt"], agg["pc_khong"]),
       "Chưa cho liên hệ → nhóm A — mục tiêu: xin số qua tin nhắn")
    kv("Phân nhóm chăm sóc",
       " · ".join("%s %d" % (k, v) for k, v in agg["do_nong"].items() if v),
       "Chấm theo 5 tín hiệu (sheet 'Chú giải ký hiệu'). Generator tự tính từ mã [①②③④⑤]")
    kv("Vấn đề quan tâm nhiều nhất",
       " · ".join("%s %d" % (k, v) for k, v in agg["van_de"][:4]),
       "Suy từ tin khách gõ tay và ảnh gửi")
    r += 1

    # ② nhóm đối tượng
    section("2 · NHÓM ĐỐI TƯỢNG THEO INSIGHT (dùng cho sheet 'Kịch bản')")
    cur = None
    for n in cfg["nhom_dt"]:
        g = n["ma"][:1]
        if g != cur:
            cur = g
            head = cfg.get("nhom_tieu_de", {}).get(g, "NHÓM %s" % g)
            col, bg = MAU_NHOM.get(g, (INK, GREY_BG))
            kv2(head, cfg.get("nhom_muc_tieu", {}).get(g, ""), bg, col, True)
        kv2("%s · %s" % (n["ma"], n["ten"]), n["insight"], WHITE, INK, True)
    r += 1

    # ③ lỗi quan sát
    section("3 · LỖI QUAN SÁT TRONG DATA — PHẢI TRÁNH (đã sửa trong kịch bản)")
    for l in cfg["loi_quan_sat"]:
        kv2("✗ " + l["loi"], l["chi_tiet"], RED_BG, RED, True)
    r += 1

    # ④ biến số chờ xác minh
    section("4 · BIẾN SỐ CHỜ XÁC MINH — ⚠ KANGNAM CHỐT TRƯỚC KHI CHẠY")
    for b in cfg["bien_so"]:
        kv2(b["ten"], b["gia_tri"], CREAM, GOLD, True)
    return ws


def sheet_chu_giai(wb, cfg):
    ws = wb.create_sheet("Chú giải ký hiệu")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", (3, 16, 78)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("B1:C1")
    c = ws.cell(1, 2, "CHÚ GIẢI KÝ HIỆU DÙNG TRONG FILE")
    c.font, c.fill, c.alignment = Font(bold=True, size=14, color=WHITE), F(NAVY), WRAP_C
    ws.row_dimensions[1].height = 28
    r = 3

    def head(txt):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(r, 2, txt)
        c.font, c.fill, c.alignment = Font(bold=True, size=11, color=WHITE), F(BLUE), WRAP_T
        r += 1

    def row(k, v, bg=None, col=INK):
        nonlocal r
        a = ws.cell(r, 2, k)
        a.font, a.alignment = Font(bold=True, size=10, color=col), WRAP_C
        if bg:
            a.fill = F(bg)
        b = ws.cell(r, 3, v)
        b.font, b.alignment = Font(size=10), WRAP_T
        r += 1

    head("CỘT 'SĐT' — khách đã để lại liên hệ chưa")
    row("Có", "Đã cho SĐT / Zalo / Viber / WhatsApp / số nước ngoài → nhóm B", GREEN_BG, GREEN)
    row("Không", "Chưa để lại liên hệ nào → nhóm A (mục tiêu: xin số)", CREAM, GOLD)
    r += 1

    head("CỘT '%s' — tình trạng khách quan tâm (suy ra từ tin và ảnh)"
         % cfg["zone"].get("nhan_van_de", "Vấn đề"))
    row("Phân loại", " · ".join(cfg["van_de"]))
    r += 1

    head("CỘT 'Ghi chú' — mã 5 tín hiệu + dấu hiệu")
    row("[①②③④⑤]", "5 chữ số 0/1 theo thứ tự 5 tín hiệu dưới đây "
                    "(vd [10110] = có ①③④, thiếu ②⑤ → 3 điểm → WARM). "
                    "Generator kiểm mã này khớp với cột 'Độ nóng', lệch là chặn build.")
    for i, t in enumerate(cfg["tin_hieu"], 1):
        row("  tín hiệu %d" % i, t)
    r += 1

    head("CỘT 'Nhóm ĐT' — nhóm đối tượng theo insight")
    for n in cfg["nhom_dt"]:
        col, bg = MAU_NHOM.get(n["ma"][:1], (INK, None))
        row(n["ma"], n["ten"], bg, col)
    r += 1

    head("CỘT 'L-code' — trạng thái vận hành (lead đang ở đâu trong luồng)")
    for l in cfg["l_code"]:
        row(l["ma"], l["nghia"])
    r += 1

    head("CỘT 'Độ nóng' — mức ưu tiên chăm sóc")
    for d in cfg["do_nong"]:
        col, bg = MAU_DO_NONG.get(d["muc"], (INK, None))
        row(d["muc"], d["nghia"], bg, col)
    r += 1

    head("KÝ HIỆU KHÁC")
    row("⚠", "Con số/thông tin CHƯA CHỐT, cần xác minh trước khi dùng "
             "(giá · % ưu đãi · ngày · hotline · số cơ sở).")
    row("[CHỜ CẬP NHẬT]", "Dữ liệu chưa có, Agent không được tự điền.")
    r += 1

    head("LƯU Ý CỐT LÕI")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ws.cell(r, 2, "L-code và HOT/WARM/COLD là HAI TRỤC KHÁC NHAU — L-code cho biết khách "
                      "đang ở BƯỚC nào trong luồng; độ nóng cho biết mức ƯU TIÊN chăm sóc. "
                      "Một khách L2.1 vẫn có thể là HOT, và một khách L4.1 vẫn có thể là COLD.")
    c.font, c.fill, c.alignment = Font(size=9, color=RED), F(CREAM), WRAP_T
    ws.row_dimensions[r].height = 42
    return ws


def sheet_phan_loai(wb, cfg, agg):
    n = agg["tong"]
    ws = wb.create_sheet("Phân loại %d HT" % n)
    ws.sheet_view.showGridLines = False
    _title(ws, "PHÂN LOẠI %d HỘI THOẠI %s — SĐT (CÓ/KHÔNG) × ĐỘ NÓNG × NHÓM ĐỐI TƯỢNG"
           % (n, cfg["zone"]["dich_vu"].upper()), 11)
    _head(ws, 2, [("STT", 5), ("Tên KH", 20), ("Ngày", 13), ("Phụ trách", 18),
                  ("Số tin", 6), ("SĐT", 7), (cfg["zone"].get("nhan_van_de", "Vấn đề"), 16), ("Nhóm ĐT", 8),
                  ("L-code", 9), ("Độ nóng", 9), ("Ghi chú (tín hiệu ①②③④⑤)", 30)])
    r = 3
    for i, h in enumerate(cfg["hoi_thoai"], 1):
        vals = [i, h["ten_kh"], h["ngay"], h["phu_trach"], h["so_tin"], h["sdt"],
                h["van_de"], h["nhom_dt"], h["l_code"], h["do_nong"], h["ghi_chu"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(r, ci, v)
            c.font, c.alignment = Font(size=9), WRAP_T
        # tô SĐT
        col, bg = (GREEN, GREEN_BG) if h["sdt"] == "Có" else (GOLD, CREAM)
        ws.cell(r, 6).font = Font(bold=True, size=9, color=col)
        ws.cell(r, 6).fill = F(bg)
        # tô độ nóng
        col, bg = MAU_DO_NONG.get(h["do_nong"], (INK, None))
        ws.cell(r, 10).font = Font(bold=True, size=9, color=col)
        if bg:
            ws.cell(r, 10).fill = F(bg)
        # tô nhóm
        col, _ = MAU_NHOM.get(h["nhom_dt"][:1], (INK, None))
        ws.cell(r, 8).font = Font(bold=True, size=9, color=col)
        r += 1
    ws.auto_filter.ref = "A2:K%d" % (r - 1)
    return ws


def sheet_kich_ban(wb, cfg):
    ws = wb.create_sheet("Kịch bản %s" % cfg["zone"]["ten"])
    ws.sheet_view.showGridLines = False
    _title(ws, "BỘ KỊCH BẢN TRỰC PAGE — %s (bám hành trình S1–S6)"
           % cfg["zone"]["dich_vu"].upper(), 10)
    _note(ws, 2, "Đây là TIN NHẮN trực page gửi khách (nhắn tin/Zalo), KHÔNG phải nội dung "
                 "landing page. Mỗi ô là một tin — gửi lần lượt, không dán cả ba. "
                 "⚠ Mọi con số (giá · % ưu đãi · ngày · hotline) lấy từ sheet 'Tổng quan' mục 4, "
                 "KHÔNG bắn giá cứng trong tin.", 10)
    _head(ws, 3, [("Nhóm", 8), ("Mã", 7), ("Giai đoạn hành trình", 18),
                  ("Insight khách hàng (nỗi lo thật)", 24), ("Trạng thái / L-code", 14),
                  ("Khi nào dùng", 20), ("Nội dung 1", 44), ("Nội dung 2", 44),
                  ("Nội dung 3", 44), ("Lưu ý an toàn", 25)])
    r = 4
    for k in cfg["kich_ban"]:
        nd = (k["noi_dung"] + ["", "", ""])[:3]
        vals = [k["nhom"], k["ma"], k["giai_doan"], k["insight"], k["l_code"],
                k["khi_dung"], nd[0], nd[1], nd[2], k["luu_y"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(r, ci, v)
            c.font, c.alignment = Font(size=9), WRAP_T
        g = k["ma"][:1]
        col, bg = MAU_NHOM.get(g, (INK, GREY_BG))
        a = ws.cell(r, 1)
        a.font, a.fill, a.alignment = Font(bold=True, size=9, color=col), F(bg), WRAP_C
        ws.cell(r, 2).font = Font(bold=True, size=11, color=NAVY)
        ws.cell(r, 2).alignment = WRAP_C
        ws.cell(r, 3).font = Font(bold=True, size=9, color=NAVY)
        ws.cell(r, 3).fill = F(L_BLUE)
        ws.cell(r, 4).font = Font(bold=True, size=9)
        ws.cell(r, 10).font = Font(size=9, color=RED)
        ws.row_dimensions[r].height = 118
        r += 1
    return ws


def sheet_cham_soc(wb, cfg, agg):
    ws = wb.create_sheet("Chăm sóc theo độ nóng")
    ws.sheet_view.showGridLines = False
    _title(ws, "CHĂM SÓC THEO ĐỘ NÓNG — HOT / WARM / COLD", 8)
    _note(ws, 2, "Độ nóng = mức ƯU TIÊN chăm sóc (quyết định SLA và bậc Telesale), "
                 "KHÁC với L-code là BƯỚC trong luồng. Số lead do generator đếm từ "
                 "sheet phân loại.", 8)
    _head(ws, 3, [("Độ nóng", 11), ("Số lead", 8), ("Nghĩa (sẵn sàng tới đâu)", 22),
                  ("Tín hiệu nhận biết", 34), ("SLA kết nối", 22), ("Nhịp chăm sóc", 34),
                  ("KB / Hành động", 26), ("Mục tiêu (KPI)", 20)])
    r = 4
    for d in cfg["do_nong"]:
        so = agg["do_nong"].get(d["muc"], 0)
        vals = [d["muc"], "%d lead" % so, d["nghia"], d["tin_hieu"], d["sla"],
                d["nhip"], d["kb"], d["kpi"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(r, ci, v)
            c.font, c.alignment = Font(size=9), WRAP_T
        col, bg = MAU_DO_NONG.get(d["muc"], (INK, GREY_BG))
        a = ws.cell(r, 1)
        a.font, a.fill, a.alignment = Font(bold=True, size=10, color=col), F(bg), WRAP_C
        ws.cell(r, 2).font = Font(bold=True, size=9, color=NAVY)
        ws.cell(r, 2).alignment = WRAP_C
        ws.cell(r, 3).font = Font(bold=True, size=9)
        ws.cell(r, 8).font = Font(bold=True, size=9, color=GREEN)
        ws.row_dimensions[r].height = 92
        r += 1
    _note(ws, r + 1, "⚠ Rủi ro 'lạm phát HOT': chấm rộng tay thì HOT tràn, Telesale bậc 1 "
                     "quá tải và SLA vỡ. Chấm đúng 5 tín hiệu, không chấm theo cảm tính. "
                     "Generator đã kiểm mã tín hiệu khớp độ nóng ở từng dòng.", 8)
    return ws


# ============================================================ TỔNG HỢP
def aggregate(cfg):
    hs = cfg["hoi_thoai"]
    n = len(hs)
    co = sum(1 for h in hs if h["sdt"] == "Có")
    dn = {}
    for d in cfg["do_nong"]:
        dn[d["muc"]] = 0
    for h in hs:
        dn[h["do_nong"]] = dn.get(h["do_nong"], 0) + 1
    vd = {}
    for h in hs:
        vd[h["van_de"]] = vd.get(h["van_de"], 0) + 1
    return {
        "tong": n, "co_sdt": co, "khong_sdt": n - co,
        "pc_co": round(co * 100.0 / n) if n else 0,
        "pc_khong": round((n - co) * 100.0 / n) if n else 0,
        "do_nong": dn,
        "van_de": sorted(vd.items(), key=lambda x: -x[1]),
    }


def build(path, outdir):
    cfg = json.load(open(path, encoding="utf-8"))
    errs = validate(cfg)
    if errs:
        print("\n⛔ TỪ CHỐI BUILD — %s" % os.path.basename(path))
        for e in errs[:40]:
            print("   • " + e)
        if len(errs) > 40:
            print("   … và %d lỗi nữa" % (len(errs) - 40))
        print("\n   Sửa JSON rồi chạy lại. KHÔNG sửa script để lách lỗi.")
        return None

    agg = aggregate(cfg)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_tong_quan(wb, cfg, agg)
    sheet_chu_giai(wb, cfg)
    sheet_phan_loai(wb, cfg, agg)
    sheet_kich_ban(wb, cfg)
    sheet_cham_soc(wb, cfg, agg)

    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, cfg["zone"]["ten_file"] + ".xlsx")
    wb.save(out)

    print("\n✅ %s" % out)
    print("   %d hội thoại · CÓ số %d (%d%%) · KHÔNG số %d (%d%%)"
          % (agg["tong"], agg["co_sdt"], agg["pc_co"], agg["khong_sdt"], agg["pc_khong"]))
    print("   Độ nóng: " + " · ".join("%s %d" % (k, v) for k, v in agg["do_nong"].items() if v))
    print("   %d nhóm đối tượng · %d kịch bản · %d L-code"
          % (len(cfg["nhom_dt"]), len(cfg["kich_ban"]), len(cfg["l_code"])))
    cho = [b["ten"] for b in cfg["bien_so"] if "⚠" in b["gia_tri"] or "CHỜ" in b["gia_tri"].upper()]
    if cho:
        print("   ⚠ CHỜ XÁC MINH: " + " · ".join(cho))
    return out


def main():
    ap = argparse.ArgumentParser(description="JSON hội thoại đã phân loại → .xlsx 5 sheet")
    ap.add_argument("pattern", help="file JSON hoặc glob, vd chats/*.json")
    ap.add_argument("-o", "--out", default="out", help="thư mục xuất (mặc định: out)")
    a = ap.parse_args()
    files = sorted(glob.glob(a.pattern))
    if not files:
        sys.exit("Không thấy file nào khớp: %s" % a.pattern)
    ok = sum(1 for f in files if build(f, a.out))
    print("\n%d/%d file build xong." % (ok, len(files)))
    if ok < len(files):
        sys.exit(1)


if __name__ == "__main__":
    main()


# ============================================================ CỔNG KIỂM TRA
# Generator dừng và in lỗi nếu:
#   • thiếu bất kỳ khối bắt buộc nào trong 11 khối;
#   • hội thoại trỏ tới nhom_dt / l_code / van_de chưa khai báo;
#   • ghi_chu không bắt đầu bằng mã 5 tín hiệu dạng [01011];
#   • ĐIỂM TÍN HIỆU KHÔNG KHỚP ĐỘ NÓNG (4–5 → HOT · 2–3 → WARM · 0–1 → COLD);
#   • nhóm A (chưa số) mà sdt='Có', hoặc nhóm B/C (đã có số) mà sdt='Không';
#   • trùng cặp (tên khách + ngày);
#   • kịch bản có nội dung chứa TỪ CẤM quảng cáo y tế;
#   • kịch bản có GIÁ CỨNG trong tin nhắn (8tr · 8.000.000 · 60 triệu…);
#   • kịch bản thiếu luu_y (rào an toàn);
#   • một nhóm đối tượng không có kịch bản nào;
#   • tin_hieu[] không đúng 5 phần tử.
#
# Hai cổng quan trọng nhất là cổng ĐIỂM↔ĐỘ NÓNG và cổng GIÁ CỨNG — đó đúng là hai
# lỗi đã quan sát được trong data thật (lạm phát HOT, và báo giá cứng trong chat).
